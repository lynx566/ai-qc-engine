from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from core.document.docx.handlers import (
    extract_table_blocks,
    extract_table_raw_rows,
    flatten_table_text,
)
from tests.extract_test_helpers import extract_docx_for_test

SAFETY_SAMPLE = (
    BACKEND_ROOT
    / "tests"
    / "fixtures"
    / "clinical_trial_doc"
    / "legacy_ib_v1_0"
    / "OOJ-VT8802-01 艾美妥珠单抗研究者手册 V1.0 2024.03.15.docx"
)


def _table_from_cells(rows: list[list[str]]):
    doc = Document()
    if not rows:
        return doc.add_table(rows=1, cols=1)
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    return table


def test_extract_table_raw_rows_simple_2x2() -> None:
    table = _table_from_cells([["A", "B"], ["1", "2"]])
    assert extract_table_raw_rows(table) == [["A", "B"], ["1", "2"]]
    assert flatten_table_text(table) == "A | B\n1 | 2"


def test_extract_table_raw_rows_preserves_empty_cells() -> None:
    table = _table_from_cells([["A", "", "C"], ["1", "2", "3"]])
    raw_rows = extract_table_raw_rows(table)
    assert raw_rows[0] == ["A", "", "C"]
    assert flatten_table_text(table) == "A | C\n1 | 2 | 3"


def test_extract_table_raw_rows_skips_all_empty_row() -> None:
    table = _table_from_cells([["A", "B"], ["", ""], ["1", "2"]])
    assert extract_table_raw_rows(table) == [["A", "B"], ["1", "2"]]


def test_extract_table_blocks_meta_includes_raw_rows() -> None:
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"

    blocks = extract_table_blocks(doc)
    assert len(blocks) == 1
    meta = blocks[0].meta or {}
    assert meta["raw_rows"] == [["A", "B"], ["1", "2"]]
    assert meta["row_count"] == 2
    assert meta["non_empty_row_count"] == 2
    assert meta["column_count"] == 2
    assert blocks[0].normalized_text == "A | B\n1 | 2"
    assert blocks[0].raw_text == "A | B\n1 | 2"


def test_safety_sample_t0001_has_raw_rows_and_flatten_text() -> None:
    bundle = extract_docx_for_test(SAFETY_SAMPLE)
    t0001 = next(block for block in bundle.table_blocks if block.block_id == "T0001")
    meta = t0001.meta or {}

    assert isinstance(meta.get("raw_rows"), list)
    assert meta["raw_rows"]
    assert all(isinstance(row, list) for row in meta["raw_rows"])
    assert all(isinstance(cell, str) for row in meta["raw_rows"] for cell in row)
    assert meta["row_count"] == 10
    assert meta["non_empty_row_count"] == len(meta["raw_rows"])
    assert meta["column_count"] == 3

    first_row = meta["raw_rows"][0]
    assert any("SOC" in cell or "不良反应" in cell for cell in first_row)
    assert t0001.normalized_text == flatten_table_text_from_meta(meta["raw_rows"])


def flatten_table_text_from_meta(raw_rows: list[list[str]]) -> str:
    """Mirror flatten_table_text output from raw_rows for regression check."""
    lines = []
    for row in raw_rows:
        cleaned = [cell for cell in row if cell]
        if cleaned:
            lines.append(" | ".join(cleaned))
    return "\n".join(lines).strip()
