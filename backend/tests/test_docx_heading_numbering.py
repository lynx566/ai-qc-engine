from __future__ import annotations

import sys
from pathlib import Path

import pytest

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from tests.fixture_builders.build_heading_numbering_fixtures import (
    EXPECTED_EFFECTIVE_HEADINGS,
    build_automatic_numbering_fixture,
    build_literal_numbering_fixture,
)
from tests.extract_test_helpers import extract_docx_for_test

SAMPLES_DIR = (
    BACKEND_ROOT
    / "tests"
    / "fixtures"
    / "clinical_trial_doc"
    / "extract"
)
AUTO_FIXTURE = SAMPLES_DIR / "ib_extract_heading_auto_numbering_001.docx"
LITERAL_FIXTURE = SAMPLES_DIR / "ib_extract_heading_literal_numbering_001.docx"
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _heading_texts(bundle) -> list[str]:
    return [
        block.normalized_text
        for block in bundle.content_blocks
        if block.kind == "heading"
    ]


def _heading_paths(bundle) -> list[list[str]]:
    return [list(block.heading_path) for block in bundle.content_blocks if block.kind == "heading"]


@pytest.fixture(scope="module")
def automatic_fixture_path() -> Path:
    if not AUTO_FIXTURE.is_file():
        build_automatic_numbering_fixture(AUTO_FIXTURE)
    return AUTO_FIXTURE


@pytest.fixture(scope="module")
def literal_fixture_path() -> Path:
    if not LITERAL_FIXTURE.is_file():
        build_literal_numbering_fixture(LITERAL_FIXTURE)
    return LITERAL_FIXTURE


def test_automatic_numbering_is_recovered(automatic_fixture_path: Path) -> None:
    doc = Document(automatic_fixture_path)
    assert [paragraph.text for paragraph in doc.paragraphs if paragraph.text] == [
        "Heading A",
        "Heading B",
        "Heading C",
        "Heading D",
    ]

    bundle = extract_docx_for_test(automatic_fixture_path)
    assert _heading_texts(bundle) == list(EXPECTED_EFFECTIVE_HEADINGS)


def test_multilevel_progression_matches_literal_fixture(
    automatic_fixture_path: Path,
    literal_fixture_path: Path,
) -> None:
    auto_bundle = extract_docx_for_test(automatic_fixture_path)
    literal_bundle = extract_docx_for_test(literal_fixture_path)

    assert _heading_texts(auto_bundle) == _heading_texts(literal_bundle)
    assert _heading_paths(auto_bundle) == _heading_paths(literal_bundle)
    assert _heading_texts(auto_bundle)[1].startswith("1.1 ")
    assert _heading_texts(auto_bundle)[3].startswith("2. ")


def test_literal_numbering_is_not_duplicated(literal_fixture_path: Path) -> None:
    bundle = extract_docx_for_test(literal_fixture_path)
    assert _heading_texts(bundle) == list(EXPECTED_EFFECTIVE_HEADINGS)


def test_non_heading_numbered_paragraph_is_untouched(tmp_path: Path) -> None:
    doc = Document()
    numbering_elm = doc.part.numbering_part._element
    numbering_elm.append(
        parse_xml(
            f"""
            <w:abstractNum xmlns:w="{_W_NS}" w:abstractNumId="902">
              <w:multiLevelType w:val="singleLevel"/>
              <w:lvl w:ilvl="0">
                <w:start w:val="1"/>
                <w:numFmt w:val="decimal"/>
                <w:lvlText w:val="%1."/>
                <w:lvlJc w:val="left"/>
              </w:lvl>
            </w:abstractNum>
            """
        )
    )
    numbering_elm.append(
        parse_xml(
            f"""
            <w:num xmlns:w="{_W_NS}" w:numId="902">
              <w:abstractNumId w:val="902"/>
            </w:num>
            """
        )
    )

    paragraph = doc.add_paragraph("Numbered body item", style="List Number")
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:numPr"))
    if existing is not None:
        p_pr.remove(existing)
    p_pr.append(
        parse_xml(
            f"""
            <w:numPr xmlns:w="{_W_NS}">
              <w:ilvl w:val="0"/>
              <w:numId w:val="902"/>
            </w:numPr>
            """
        )
    )

    path = tmp_path / "numbered_list_item.docx"
    doc.save(path)

    bundle = extract_docx_for_test(path)
    blocks = [block for block in bundle.content_blocks if block.normalized_text]
    assert len(blocks) == 1
    assert blocks[0].kind == "list_item"
    assert blocks[0].normalized_text == "Numbered body item"


def test_normalization_audit_meta_unchanged_for_automatic_heading(
    automatic_fixture_path: Path,
) -> None:
    bundle = extract_docx_for_test(automatic_fixture_path)
    headings = [block for block in bundle.content_blocks if block.kind == "heading"]
    assert headings
    for heading in headings:
        meta = heading.meta or {}
        assert meta.get("raw_text_changed_by_normalize") is False
        assert "raw_text_before_normalize" not in meta


def test_existing_literal_heading_fixture_still_passes_extract(
    literal_fixture_path: Path,
) -> None:
    bundle = extract_docx_for_test(literal_fixture_path)
    assert len(_heading_texts(bundle)) == 4
    assert all("Heading" in text for text in _heading_texts(bundle))
