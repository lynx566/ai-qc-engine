"""Targeted tests for DOCX block-level presentation metadata extraction."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from core.document.docx.presentation import (
    extract_block_presentation_meta,
    infer_toc_level_from_style_name,
)
from tests.extract_test_helpers import extract_docx_for_test
from tests.test_docx_revision_extract import (
    MINIMAL_FIXTURE,
    _append_ins,
    _extract_bundle,
)


def _append_del_bold(parent, text: str, *, revision_id: str = "99") -> None:
    del_el = OxmlElement("w:del")
    del_el.set(qn("w:id"), revision_id)
    del_el.set(qn("w:author"), "test")
    del_el.set(qn("w:date"), "2026-01-01T00:00:00Z")
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_properties.append(OxmlElement("w:b"))
    run.append(run_properties)
    del_text = OxmlElement("w:delText")
    del_text.text = text
    run.append(del_text)
    del_el.append(run)
    parent.append(del_el)


TOC_EXTRACT_FIXTURE = (
    BACKEND_ROOT
    / "tests"
    / "fixtures"
    / "clinical_trial_doc"
    / "extract"
    / "frontend_reader_canonical_showcase_sample.docx"
)


def _presentation(block) -> dict:
    return block.meta.get("presentation", {})


def test_cover_style_paragraph_presentation() -> None:
    doc = Document()
    paragraph = doc.add_paragraph("Cover Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(36)
    run = paragraph.runs[0]
    run.bold = True
    run.font.size = Pt(24)

    bundle = _extract_bundle(doc)
    block = bundle.content_blocks[0]
    presentation = _presentation(block)

    assert presentation["alignment"] == "center"
    assert presentation["font_size_pt"] == 24.0
    assert presentation["bold"] is True
    assert presentation["space_before_pt"] == 36.0


def test_body_paragraph_spacing_and_indent() -> None:
    doc = Document()
    paragraph = doc.add_paragraph("正文首段内容。")
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = 1.45
    paragraph.runs[0].font.size = Pt(10.5)

    bundle = _extract_bundle(doc)
    presentation = _presentation(bundle.content_blocks[0])

    assert presentation["indent_first_line_pt"] == 21.0
    assert presentation["line_spacing_multiple"] == 1.45
    assert presentation["space_after_pt"] == 6.0
    assert presentation["font_size_pt"] == 10.5


def test_mixed_inline_bold_omits_block_level_bold() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("总体安全性维持可接受，其中")
    bold_run = paragraph.add_run("重要风险")
    bold_run.bold = True
    paragraph.add_run("仅局部 bold。")

    bundle = _extract_bundle(doc)
    presentation = _presentation(bundle.content_blocks[0])

    assert "bold" not in presentation


def test_uniform_italic_is_persisted() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("整段均为 italic 的正文。")
    run.italic = True

    bundle = _extract_bundle(doc)
    presentation = _presentation(bundle.content_blocks[0])

    assert presentation["italic"] is True


def test_deleted_run_formatting_does_not_pollute_visible_presentation() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("KEEP ")
    _append_del_bold(paragraph._p, "OLD")
    paragraph.add_run("NEW")

    bundle = _extract_bundle(doc)
    block = bundle.content_blocks[0]
    presentation = _presentation(block)

    assert block.raw_text == "KEEP NEW"
    assert "bold" not in presentation


def test_revision_replacement_keeps_inserted_visible_formatting_only() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("prefix ")
    _append_del_bold(paragraph._p, "oldword", revision_id="4")
    _append_ins(paragraph._p, "newword", revision_id="5")

    bundle = _extract_bundle(doc)
    block = bundle.content_blocks[0]

    assert block.raw_text == "prefix newword"
    assert "bold" not in _presentation(block)


def test_toc_level_from_style_name_parser() -> None:
    assert infer_toc_level_from_style_name("TOC 1") == 1
    assert infer_toc_level_from_style_name("toc 2") == 2
    assert infer_toc_level_from_style_name("TOC 11") == 1
    assert infer_toc_level_from_style_name("TOC 21") == 2
    assert infer_toc_level_from_style_name("TOC Heading") is None


def test_toc_blocks_persist_toc_level_on_extract() -> None:
    bundle = extract_docx_for_test(TOC_EXTRACT_FIXTURE)
    toc_blocks = [block for block in bundle.front_matter_blocks if block.kind == "toc"]
    assert toc_blocks

    leveled = [
        block
        for block in toc_blocks
        if block.meta.get("presentation", {}).get("toc_level") is not None
    ]

    assert leveled


def test_toc_style_assignment_persists_level_without_text_heuristic() -> None:
    doc = Document()
    paragraph = doc.add_paragraph("章节\t3")
    try:
        paragraph.style = doc.styles["TOC 2"]
    except KeyError:
        presentation = extract_block_presentation_meta(
            paragraph,
            style_name="TOC 2",
            kind="toc",
        )
        assert presentation["toc_level"] == 2
        return

    bundle = _extract_bundle(doc)
    toc_block = bundle.front_matter_blocks[0]
    assert toc_block.kind == "toc"
    assert _presentation(toc_block)["toc_level"] == 2


def test_ordinary_document_text_contract_unchanged() -> None:
    baseline = extract_docx_for_test(MINIMAL_FIXTURE)

    def snapshot(bundle):
        blocks = (
            bundle.front_matter_blocks
            + bundle.content_blocks
            + bundle.table_blocks
            + bundle.format_blocks
        )
        return [
            {
                "block_id": block.block_id,
                "kind": block.kind,
                "raw_text": block.raw_text,
                "normalized_text": block.normalized_text,
                "sentence_ids": [sentence.sentence_id for sentence in block.sentences],
            }
            for block in blocks
        ]

    current = extract_docx_for_test(MINIMAL_FIXTURE)
    assert snapshot(baseline) == snapshot(current)
