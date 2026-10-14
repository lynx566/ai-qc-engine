"""Build literal vs automatic heading-numbering extract fixtures."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
OUTPUT_DIR = (
    Path(__file__).resolve().parents[1]
    / "fixtures"
    / "clinical_trial_doc"
    / "extract"
)

HEADING_SEQUENCE: tuple[tuple[int, str], ...] = (
    (1, "Heading A"),
    (2, "Heading B"),
    (2, "Heading C"),
    (1, "Heading D"),
)

LITERAL_HEADING_TEXT: tuple[tuple[int, str], ...] = (
    (1, "1. Heading A"),
    (2, "1.1 Heading B"),
    (2, "1.2 Heading C"),
    (1, "2. Heading D"),
)

AUTO_HEADING_TEXT: tuple[tuple[int, str], ...] = HEADING_SEQUENCE
EXPECTED_EFFECTIVE_HEADINGS: tuple[str, ...] = tuple(
    text for _, text in LITERAL_HEADING_TEXT
)


def _ensure_multilevel_numbering(doc: Document, *, abstract_id: int, num_id: int) -> None:
    numbering_elm = doc.part.numbering_part._element
    has_abstract = any(
        int(abstract.get(qn("w:abstractNumId"))) == abstract_id
        for abstract in numbering_elm.findall(qn("w:abstractNum"))
    )
    if not has_abstract:
        numbering_elm.append(
            parse_xml(
                f"""
                <w:abstractNum xmlns:w="{_W_NS}" w:abstractNumId="{abstract_id}">
                  <w:multiLevelType w:val="multilevel"/>
                  <w:lvl w:ilvl="0">
                    <w:start w:val="1"/>
                    <w:numFmt w:val="decimal"/>
                    <w:lvlText w:val="%1."/>
                    <w:lvlJc w:val="left"/>
                  </w:lvl>
                  <w:lvl w:ilvl="1">
                    <w:start w:val="1"/>
                    <w:numFmt w:val="decimal"/>
                    <w:lvlText w:val="%1.%2"/>
                    <w:lvlJc w:val="left"/>
                  </w:lvl>
                </w:abstractNum>
                """
            )
        )
    has_num = any(
        int(num.get(qn("w:numId"))) == num_id for num in numbering_elm.findall(qn("w:num"))
    )
    if not has_num:
        numbering_elm.append(
            parse_xml(
                f"""
                <w:num xmlns:w="{_W_NS}" w:numId="{num_id}">
                  <w:abstractNumId w:val="{abstract_id}"/>
                </w:num>
                """
            )
        )


def _set_style_numbering(doc: Document, *, num_id: int) -> None:
    for level in (1, 2):
        style = doc.styles[f"Heading {level}"]
        element = style.element
        p_pr = element.pPr
        if p_pr is None:
            p_pr = OxmlElement("w:pPr")
            element.insert(0, p_pr)
        existing = p_pr.find(qn("w:numPr"))
        if existing is not None:
            p_pr.remove(existing)
        p_pr.append(
            parse_xml(
                f"""
                <w:numPr xmlns:w="{_W_NS}">
                  <w:ilvl w:val="{level - 1}"/>
                  <w:numId w:val="{num_id}"/>
                </w:numPr>
                """
            )
        )


def build_literal_numbering_fixture(path: Path | None = None) -> Path:
    output = path or OUTPUT_DIR / "ib_extract_heading_literal_numbering_001.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for level, heading_text in LITERAL_HEADING_TEXT:
        doc.add_heading(heading_text, level=level)
    doc.save(output)
    return output


def build_automatic_numbering_fixture(path: Path | None = None) -> Path:
    output = path or OUTPUT_DIR / "ib_extract_heading_auto_numbering_001.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _ensure_multilevel_numbering(doc, abstract_id=901, num_id=901)
    _set_style_numbering(doc, num_id=901)
    for level, heading_text in HEADING_SEQUENCE:
        doc.add_heading(heading_text, level=level)
    doc.save(output)
    return output


def main() -> None:
    literal = build_literal_numbering_fixture()
    automatic = build_automatic_numbering_fixture()
    print(literal)
    print(automatic)


if __name__ == "__main__":
    main()
