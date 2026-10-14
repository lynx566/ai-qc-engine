"""Targeted regression tests for DOCX Final / No Markup revision-aware Extract."""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from core.document.docx.revision_text import (
    extract_paragraph_revision_aware_text,
    revision_aware_text_matches_legacy_paragraph,
)
from tests.extract_test_helpers import extract_docx_for_test

MINIMAL_FIXTURE = (
    BACKEND_ROOT
    / "tests"
    / "fixtures"
    / "clinical_trial_doc"
    / "extract"
    / "ib_extract_minimal_001.docx"
)


def _append_run(parent, text: str) -> None:
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    parent.append(run)


def _append_ins(parent, text: str, *, revision_id: str = "1") -> None:
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), revision_id)
    ins.set(qn("w:author"), "test")
    ins.set(qn("w:date"), "2026-01-01T00:00:00Z")
    _append_run(ins, text)
    parent.append(ins)


def _append_del(parent, text: str, *, revision_id: str = "2") -> None:
    del_el = OxmlElement("w:del")
    del_el.set(qn("w:id"), revision_id)
    del_el.set(qn("w:author"), "test")
    del_el.set(qn("w:date"), "2026-01-01T00:00:00Z")
    run = OxmlElement("w:r")
    del_text = OxmlElement("w:delText")
    del_text.text = text
    run.append(del_text)
    del_el.append(run)
    parent.append(del_el)


def _append_move_from(parent, text: str, *, revision_id: str = "3") -> None:
    move_from = OxmlElement("w:moveFrom")
    move_from.set(qn("w:id"), revision_id)
    move_from.set(qn("w:author"), "test")
    move_from.set(qn("w:date"), "2026-01-01T00:00:00Z")
    _append_run(move_from, text)
    parent.append(move_from)


def _append_move_to(parent, text: str, *, revision_id: str = "3") -> None:
    move_to = OxmlElement("w:moveTo")
    move_to.set(qn("w:id"), revision_id)
    move_to.set(qn("w:author"), "test")
    move_to.set(qn("w:date"), "2026-01-01T00:00:00Z")
    _append_run(move_to, text)
    parent.append(move_to)


def _extract_bundle(doc: Document):
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
        doc.save(handle.name)
        return extract_docx_for_test(handle.name)


def test_clean_doc_matches_legacy_paragraph_text() -> None:
    bundle = extract_docx_for_test(MINIMAL_FIXTURE)
    assert bundle.content_blocks
    assert bundle.document_meta.get("parse_warnings") == []

    doc = Document(MINIMAL_FIXTURE)
    for paragraph in doc.paragraphs:
        assert revision_aware_text_matches_legacy_paragraph(paragraph)


def test_insertion_only_paragraph_is_extracted() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    _append_ins(paragraph._p, "INSERTED_ONLY")

    bundle = _extract_bundle(doc)
    assert len(bundle.content_blocks) == 1
    assert bundle.content_blocks[0].raw_text == "INSERTED_ONLY"
    assert "tracked_insertions_detected" in bundle.document_meta["parse_warnings"]


def test_deleted_text_is_excluded_from_final_text() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("KEEP ")
    _append_del(paragraph._p, "DROP")
    paragraph.add_run(" TAIL")

    raw_text, _obs = extract_paragraph_revision_aware_text(paragraph)
    assert raw_text == "KEEP  TAIL"

    bundle = _extract_bundle(doc)
    assert bundle.content_blocks[0].raw_text == "KEEP  TAIL"
    assert "tracked_deletions_detected" in bundle.document_meta["parse_warnings"]


def test_replacement_pattern_uses_inserted_word_only() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("prefix ")
    _append_del(paragraph._p, "oldword")
    _append_ins(paragraph._p, "newword", revision_id="4")
    paragraph.add_run(" suffix")

    bundle = _extract_bundle(doc)
    assert bundle.content_blocks[0].raw_text == "prefix newword suffix"
    warnings = bundle.document_meta["parse_warnings"]
    assert "tracked_insertions_detected" in warnings
    assert "tracked_deletions_detected" in warnings


def test_table_cell_revision_semantics_match_paragraph() -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.add_run("CELL_")
    _append_ins(paragraph._p, "INS", revision_id="5")
    _append_del(paragraph._p, "DEL", revision_id="6")

    bundle = _extract_bundle(doc)
    assert bundle.table_blocks[0].raw_text == "CELL_INS"


def test_comment_marked_body_text_is_unchanged() -> None:
    doc = Document()
    paragraph = doc.add_paragraph("COMMENTED_BODY_TEXT")
    paragraph.runs[0].mark_comment_range(paragraph.runs[0], 0)

    bundle = _extract_bundle(doc)
    assert bundle.content_blocks[0].raw_text == "COMMENTED_BODY_TEXT"
    assert bundle.document_meta.get("parse_warnings") == []


def test_move_revision_includes_move_to_and_excludes_move_from() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("BASE ")
    _append_move_from(paragraph._p, "MOVED_OUT")
    _append_move_to(paragraph._p, "MOVED_IN")

    bundle = _extract_bundle(doc)
    assert bundle.content_blocks[0].raw_text == "BASE MOVED_IN"
    assert "tracked_moves_detected" in bundle.document_meta["parse_warnings"]


def test_revision_identity_closure_assigns_sentences() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("临床安全性。")
    _append_ins(paragraph._p, "不良事件。", revision_id="7")

    bundle = _extract_bundle(doc)
    block = bundle.content_blocks[0]
    assert block.block_id == "C0001"
    assert block.sentences
    assert block.sentences[0].sentence_id == "C0001:S001"
    assert "不良事件" in block.raw_text
