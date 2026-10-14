"""Regression tests for Extract Block.raw_text provenance contract."""

from __future__ import annotations

import json
import sys
from dataclasses import asdict
from pathlib import Path

import pytest
from docx import Document

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from artifacts.privacy_safety import (
    PUBLIC_ANALYSIS_BUNDLE_FORBIDDEN_KEYS,
    assert_no_forbidden_keys,
)
from core.document.document_schema import BUNDLE_SCHEMA_VERSION, Block, SourceRef
from core.document.docx.handlers import join_paragraph_texts
from core.document.sentence_index import apply_sentence_index_to_block
from core.normalization.text_normalize import normalize_text
from tests.fixture_builders.build_heading_numbering_fixtures import (
    build_automatic_numbering_fixture,
    build_literal_numbering_fixture,
)
from tests.clinical_program_fixtures import (
    canonical_clinical_program_fields,
    prepare_analysis_bundle_for_test,
)
from tests.extract_test_helpers import extract_docx_for_test

SAMPLES_DIR = (
    BACKEND_ROOT
    / "tests"
    / "fixtures"
    / "clinical_trial_doc"
    / "extract"
)
AUTO_FIXTURE = SAMPLES_DIR / "ib_extract_heading_auto_numbering_001.docx"
LITERAL_FIXTURE = SAMPLES_DIR / "ib_extract_heading_literal_numbering_001.docx"


def _ensure_auto_fixture() -> Path:
    if not AUTO_FIXTURE.is_file():
        build_automatic_numbering_fixture(AUTO_FIXTURE)
    return AUTO_FIXTURE


def _ensure_literal_fixture() -> Path:
    if not LITERAL_FIXTURE.is_file():
        build_literal_numbering_fixture(LITERAL_FIXTURE)
    return LITERAL_FIXTURE


def _bundle_with_raw_text(*, raw_text: str, normalized_text: str) -> dict:
    return {
        "doc_id": "doc_raw_text_safety",
        "file_hash": "sha256:" + "b" * 64,
        "file_name": "raw_text.docx",
        "source_type": "docx",
        "document_meta": {"language": "zh", "bundle_schema_version": "0.2.0"},
        "front_matter_blocks": [],
        "content_blocks": [
            {
                "block_id": "C0001",
                "kind": "paragraph",
                "raw_text": raw_text,
                "normalized_text": normalized_text,
                "order": 1,
                "doc_order_index": 0,
                "source_ref": {
                    "source_type": "docx",
                    "locator_type": "paragraph",
                    "locator_payload": {"paragraph_index": 0},
                },
                "sentences": [
                    {
                        "sentence_id": "C0001:S001",
                        "raw_text": raw_text,
                        "normalized_text": normalized_text,
                        "start": 0,
                        "end": len(normalized_text),
                    }
                ],
            }
        ],
        "table_blocks": [],
        "format_blocks": [],
    }


def test_paragraph_raw_text_captures_pre_normalization(tmp_path: Path) -> None:
    visible = "表2\u00a0\u200b   OOJ\u00adVT8802-02"
    doc = Document()
    doc.add_paragraph(visible)
    path = tmp_path / "normalize_changed.docx"
    doc.save(path)

    bundle = extract_docx_for_test(path)
    block = bundle.content_blocks[0]
    assert block.raw_text == visible
    assert block.normalized_text == normalize_text(visible)
    assert block.raw_text != block.normalized_text


def test_paragraph_raw_text_present_when_unchanged(tmp_path: Path) -> None:
    visible = "摘要"
    doc = Document()
    doc.add_paragraph(visible)
    path = tmp_path / "unchanged.docx"
    doc.save(path)

    bundle = extract_docx_for_test(path)
    block = bundle.content_blocks[0]
    assert block.raw_text == visible
    assert block.normalized_text == visible


def test_heading_auto_number_raw_text_excludes_derived_prefix() -> None:
    path = _ensure_auto_fixture()
    bundle = extract_docx_for_test(path)
    headings = [block for block in bundle.content_blocks if block.kind == "heading"]
    assert headings
    second = headings[1]
    assert second.raw_text == "Heading B"
    assert second.normalized_text == "1.1 Heading B"


def test_heading_literal_number_in_raw_text() -> None:
    path = _ensure_literal_fixture()
    bundle = extract_docx_for_test(path)
    headings = [block for block in bundle.content_blocks if block.kind == "heading"]
    assert headings
    for heading in headings:
        assert heading.raw_text == heading.normalized_text
        assert heading.normalized_text


def test_table_raw_text_differs_when_normalization_changes(tmp_path: Path) -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A\u00a0B"
    table.cell(0, 1).text = "1\u200b2"
    path = tmp_path / "table_raw.docx"
    doc.save(path)

    bundle = extract_docx_for_test(path)
    block = bundle.table_blocks[0]
    assert block.raw_text != block.normalized_text
    assert block.normalized_text == "A B | 12"
    assert " | " in block.raw_text
    assert block.raw_text.count(" | ") == block.normalized_text.count(" | ")
    assert block.raw_text.count("\n") == block.normalized_text.count("\n")


def test_header_footer_raw_text_aggregation(tmp_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Header\u00a0A"
    new_para = section.header.add_paragraph("Header B")
    new_para.text = "Header B"
    path = tmp_path / "header.docx"
    doc.save(path)

    bundle = extract_docx_for_test(path)
    header = next(block for block in bundle.format_blocks if block.kind == "header")
    assert header.raw_text == "Header\u00a0A\nHeader B"
    assert header.normalized_text == "Header A\nHeader B"


def test_join_paragraph_texts_matches_format_block_gate() -> None:
    doc = Document()
    section = doc.sections[0]
    section.footer.paragraphs[0].text = "   "
    section.footer.add_paragraph("Footer\u00a01")
    raw_text, text = join_paragraph_texts(section.footer.paragraphs)
    assert raw_text == "Footer\u00a01"
    assert text == "Footer 1"


def test_persisted_bundle_contains_raw_text(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Persist raw_text")
    path = tmp_path / "persist.docx"
    doc.save(path)

    bundle = extract_docx_for_test(path)
    payload = asdict(bundle)
    assert payload["document_meta"]["bundle_schema_version"] == BUNDLE_SCHEMA_VERSION == "0.2.1"
    for group in (
        "front_matter_blocks",
        "content_blocks",
        "table_blocks",
        "format_blocks",
    ):
        for block in payload[group]:
            assert "raw_text" in block
            assert isinstance(block["raw_text"], str)

    serialized = json.dumps(payload, ensure_ascii=False)
    assert '"raw_text"' in serialized


def test_block_dataclass_allows_empty_raw_text() -> None:
    block = Block(
        block_id="C0001",
        kind="paragraph",
        raw_text="",
        normalized_text="",
        loc="BODY:P1",
        source_ref=SourceRef(
            source_type="docx",
            locator_type="paragraph",
            locator_payload={"paragraph_index": 0},
        ),
    )
    assert block.raw_text == ""


def test_prepare_identity_path_strips_sentence_raw_text() -> None:
    bundle = _bundle_with_raw_text(
        raw_text="visible raw",
        normalized_text="visible raw",
    )
    result = prepare_analysis_bundle_for_test(
        bundle,
        profile=None,
        language="zh",
        **canonical_clinical_program_fields(),
    )
    sentence = result.analysis_bundle["content_blocks"][0]["sentences"][0]
    assert "raw_text" not in sentence
    assert sentence["normalized_text"] == "visible raw"


def test_prepare_identity_path_strips_block_raw_text() -> None:
    bundle = _bundle_with_raw_text(
        raw_text="visible raw",
        normalized_text="visible raw",
    )
    result = prepare_analysis_bundle_for_test(
        bundle,
        profile=None,
        language="zh",
        **canonical_clinical_program_fields(),
    )
    block = result.analysis_bundle["content_blocks"][0]
    assert "raw_text" not in block
    assert block["analysis_text"] == "visible raw"


def test_prepare_privacy_on_public_strips_raw_text() -> None:
    from privacy.config import PrivacyConfig
    from privacy.transform import apply_privacy_transform
    from plugins.clinical_trial_doc.privacy.profile import (
        CLINICAL_TRIAL_DOC_PRIVACY_PROFILE,
    )

    bundle = _bundle_with_raw_text(
        raw_text="VT8802 raw visible",
        normalized_text="VT8802 raw visible",
    )
    config = PrivacyConfig.default().with_transform_detectors(
        "after_extract",
        ["registry"],
    )
    result = apply_privacy_transform(
        bundle,
        "after_extract",
        config,
        profile=CLINICAL_TRIAL_DOC_PRIVACY_PROFILE,
        language="zh",
        clinical_program_id=canonical_clinical_program_fields()["clinical_program_id"],
    )
    assert result.analysis_bundle["privacy_enabled"] is True
    for group in (
        "content_blocks",
        "table_blocks",
        "format_blocks",
        "front_matter_blocks",
    ):
        for block in result.analysis_bundle.get(group) or []:
            assert "raw_text" not in block


def test_public_analysis_bundle_forbids_raw_text_key() -> None:
    payload = {
        "content_blocks": [{"block_id": "C1", "raw_text": "leak"}],
    }
    with pytest.raises(ValueError, match="raw_text"):
        assert_no_forbidden_keys(
            payload,
            PUBLIC_ANALYSIS_BUNDLE_FORBIDDEN_KEYS,
            label="analysis_bundle",
        )


def test_block_manifest_does_not_carry_raw_text() -> None:
    from tests.block_manifest_test_helpers import build_block_manifest

    analysis_bundle = {
        "analysis_bundle_id": "prepare_analysis_raw_text_contract",
        "doc_id": "doc_raw_text_contract",
        "file_hash": "sha256:" + "c" * 64,
        "content_blocks": [
            {
                "block_id": "C0001",
                "kind": "heading",
                "order": 1,
                "analysis_text": "masked heading",
                "raw_text": "must not leak",
            }
        ],
        "table_blocks": [],
    }
    manifest = build_block_manifest(analysis_bundle)
    for block in (manifest.get("blocks") or []) + (manifest.get("table_blocks") or []):
        assert "raw_text" not in block


def test_sentence_index_unchanged_after_raw_text_added(tmp_path: Path) -> None:
    text = "第一句。第二句！"
    block = Block(
        block_id="C0099",
        kind="paragraph",
        raw_text=text + "\n",
        normalized_text=text,
        loc="BODY:P1",
        source_ref=SourceRef(
            source_type="docx",
            locator_type="paragraph",
            locator_payload={"paragraph_index": 0},
        ),
    )
    apply_sentence_index_to_block(block)
    assert block.normalized_text == text
    assert len(block.sentences) == 2
    assert block.sentences[0].sentence_id == "C0099:S001"
    assert block.sentences[0].normalized_text == "第一句。"
